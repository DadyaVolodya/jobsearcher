#!/usr/bin/env python3
"""Сборка презентации ФСП: слайды 7–11 по образцу «хулиганы», пустые после 11 удаляются."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.slide import Slide

SRC = Path("/Users/sparrow/Downloads/ФСП 2/ЛЦТ_2026_Шаблон презентации ФСП.pptx")
OUT_DIR = Path(__file__).resolve().parents[1] / "docs" / "presentation"
OUT_PPTX = OUT_DIR / "FSP_Jobsearcher_presentation.pptx"
OUT_PDF = OUT_DIR / "FSP_Jobsearcher_presentation.pdf"


def set_text(shape, text: str) -> None:
    """Replace text keeping first-run formatting; reuse existing paragraphs."""
    if not shape.has_text_frame:
        return
    lines = text.split("\n") if text else [""]
    tf = shape.text_frame
    # Ensure enough paragraphs
    while len(tf.paragraphs) < len(lines):
        tf.add_paragraph()
    for i, para in enumerate(tf.paragraphs):
        value = lines[i] if i < len(lines) else ""
        if para.runs:
            para.runs[0].text = value
            for run in para.runs[1:]:
                run.text = ""
        else:
            para.text = value
        if i > 0 and i < len(lines) and tf.paragraphs[0].runs:
            # copy basic font from first paragraph when we added empty paras
            src = tf.paragraphs[0].runs[0].font
            if para.runs:
                if src.size:
                    para.runs[0].font.size = src.size
                if src.name:
                    para.runs[0].font.name = src.name


def delete_slide(prs: Presentation, index: int) -> None:
    sld_id = prs.slides._sldIdLst[index]
    r_id = sld_id.rId
    prs.part.drop_rel(r_id)
    prs.slides._sldIdLst.remove(sld_id)


def fill_slide7(slide: Slide) -> None:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.name.startswith("Заголовок"):
            set_text(sh, "ФСП Карьера")
        elif "Цифровая платформа" in sh.text_frame.text or sh.name.startswith("Текст"):
            if sh.text_frame.text.strip() or "Цифровая" in sh.text_frame.text:
                set_text(
                    sh,
                    "Платформа подбора ИТ-специалистов\n"
                    "с подтверждённым грейдом и приоритетом достижений ФСП",
                )


def fill_slide8(slide: Slide) -> None:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("Капитан:"):
            set_text(
                sh,
                "Капитан: Воробьев А. В., бэкенд-разработчик\n"
                "Кол-во участников: 2 человека\n"
                "Краткое описание: команда сложилась вокруг общих интересов\n"
                "и стремления довести практическую задачу до рабочего результата.\n"
                "Города: Пенза и Тула",
            )
        elif t.strip() == "В чем суть вашего решения":
            set_text(
                sh,
                "Кандидат проходит анкету и тест, получает категорию и грейд.\n"
                "Работодатель ищет по подтверждённым навыкам и сам направляет\n"
                "приглашение с зарплатой до начала чата.",
            )
        elif t.strip() == "Что делает ваше решение уникальным или инновационным?":
            set_text(
                sh,
                "Обратная механика подбора: адресное приглашение, а не отклик «вслепую».\n"
                "Грейд только тестом; достижения ФСП поднимают в выдаче.\n"
                "Контакты скрыты до принятия предложения.",
            )
        elif sh.name.startswith("Заголовок"):
            set_text(sh, "ФСП Карьера")


def fill_slide9(slide: Slide) -> None:
    members = [
        (
            "Воробьев Александр",
            "Капитан, бэкенд-разработчик\n@AleksandrCaptainSparrow\n+79513583328\nг. Пенза",
        ),
        (
            "Егоров Дмитрий",
            "Frontend-разработчик\n@Egorov00\n+79066284052\nг. Тула",
        ),
    ]
    name_boxes = []
    role_boxes = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text.strip()
        if txt == "Имя Фамилия":
            name_boxes.append(sh)
        elif txt.startswith("Роль в команде"):
            role_boxes.append(sh)
        elif sh.name.startswith("Заголовок"):
            set_text(sh, "Команда «ФСП Карьера»")
    name_boxes.sort(key=lambda s: s.left)
    role_boxes.sort(key=lambda s: s.left)
    for i, sh in enumerate(name_boxes):
        if i < len(members):
            set_text(sh, members[i][0])
        else:
            set_text(sh, "")
    for i, sh in enumerate(role_boxes):
        if i < len(members):
            set_text(sh, members[i][1])
        else:
            set_text(sh, "")


def fill_slide10(slide: Slide) -> None:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if sh.name.startswith("Заголовок"):
            set_text(sh, "О команде и пути к решению")
        elif "как вы собрались" in t:
            set_text(
                sh,
                "Мы объединились благодаря общим интересам и желанию решить\n"
                "практическую задачу спец. трека ФСП. В команде дополняют друг друга\n"
                "компетенции в backend-сервисах и фронтенде личных кабинетов.",
            )
        elif "Почему вы выбрали именно эту задачу" in t:
            set_text(sh, "Почему выбрали эту задачу")
        elif "Что вас вдохновило" in t:
            set_text(
                sh,
                "Задача даёт прикладной результат для рынка труда и сочетает\n"
                "тестирование навыков с объяснимым подбором. Можно опереться\n"
                "на сильные стороны каждого участника.",
            )
        elif "сложностями или вызовами" in t:
            set_text(sh, "Сложности и вызовы")
        elif "самых интересных или сложных моментах" in t:
            set_text(
                sh,
                "Основные вызовы - антисписывание в тестах, обязательная зарплата\n"
                "в приглашении без падений API, согласия 152-ФЗ и непрочитанные\n"
                "сообщения. Справились за счёт чёткого разделения ролей\n"
                "и фиксации решений в контракте API.",
            )


def fill_slide11(slide: Slide) -> None:
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if t.startswith("Опишите в чем техническая"):
            set_text(
                sh,
                "Java 21 + Spring Boot, PostgreSQL, Flyway, OpenAPI, Docker Compose.\n"
                "Конвейер: анкета → тест A/B/C → categoryCode → matching с explain и бустом ФСП.\n"
                "Приглашения salary-first, чаты, unread; кабинеты candidate / employer (React).",
            )
        elif t.startswith("Опишите ваши идеи по дальнейшему"):
            set_text(
                sh,
                "Платформа для работодателей и участников ФСП: подтверждённый грейд до контакта.\n"
                "Дальше: боевой реестр ФСП, Keycloak, расширенный банк заданий, PDF-профиль.\n"
                "Развёртывание: docker compose -f docker-compose.demo.yml up -d --build",
            )


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"Нет шаблона: {SRC}")
    prs = Presentation(str(SRC))
    fill_slide7(prs.slides[6])
    fill_slide8(prs.slides[7])
    fill_slide9(prs.slides[8])
    fill_slide10(prs.slides[9])
    fill_slide11(prs.slides[10])

    # удалить всё после 11-го слайда (пустые шаблоны / иконки)
    while len(prs.slides) > 11:
        delete_slide(prs, len(prs.slides) - 1)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"Wrote {OUT_PPTX} ({len(prs.slides)} slides)")

    # verify
    check = Presentation(str(OUT_PPTX))
    assert len(check.slides) == 11, len(check.slides)
    for i, s in enumerate(check.slides, 1):
        texts = [sh.text_frame.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        print(f"slide {i}: {len(texts)} text boxes, sample={texts[0][:60]!r}" if texts else f"slide {i}: empty?")


if __name__ == "__main__":
    main()
